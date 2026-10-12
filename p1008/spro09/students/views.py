
from django.shortcuts import render, redirect
from students.models import Stu
from pathlib import Path
from django.conf import settings
from openpyxl import load_workbook


def swrite(request):
    if request.method == 'GET':
        print("GET 페이지가 로딩되었습니다.")
        return render(request, 'swrite.html')
    elif request.method == 'POST':
        print("POST 페이지가 로딩")
        name = request.POST.get('name')
        major = request.POST.get('major')
        grade = request.POST.get('grade')
        age = request.POST.get('age')
        gender = request.POST.get('gender')

        # 기존 학생정보 저장 코드 유지
        qs = Stu(name=name, major=major, grade=grade, age=age, gender=gender)
        qs.save()
        print(name, major, grade, age, gender)
        return redirect('/')


def slist(request):
    # 프로젝트 최상위 폴더의 students.xlsx 경로
    excel_path = Path(settings.BASE_DIR) / 'students.xlsx'

    # 엑셀 파일이 존재할 때만 데이터 읽기
    if not excel_path.exists():
        return render(request, 'slist.html', {
            'headers': [],
            'rows': [],
            'error': 'students.xlsx 파일을 프로젝트 최상위 폴더에 넣어야 함.'
        })

    # 기존 students.xlsx의 첫 번째 시트 읽기
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        data = list(sheet.iter_rows(values_only=True))

        # 첫 번째 행은 기존 엑셀의 열 제목
        headers = list(data[0]) if data else []

        # 나머지 행은 학생 데이터
        rows = [
            list(row)
            for row in data[1:]
            if any(value is not None for value in row)
        ]
    finally:
        # 엑셀 파일 사용 종료
        workbook.close()

    return render(request, 'slist.html', {
        'headers': headers,
        'rows': rows,
        'error': ''
    })
