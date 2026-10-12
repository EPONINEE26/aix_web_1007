"""students.xlsx 데이터를 stuscore StudentScore 모델로 가져오는 관리 명령"""

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import connection
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from stuscore.models import StudentScore


# manage.py와 students.xlsx가 함께 있는 프로젝트 최상위 폴더
EXCEL_FILE = Path(__file__).resolve().parents[3] / 'students.xlsx'


# 엑셀 헤더와 모델 필드 이름을 연결
COLUMN_ALIASES = {
    '번호': 'no',
    '이름': 'name',
    '학교': 'school',
    '학년': 'grade',
    '키': 'stature',
    '국어': 'kor',
    '영어': 'eng',
    '수학': 'math',
    'SW특기': 'sw',
}


def _text(value):
    """None은 빈 문자열로 바꾸고 나머지는 공백을 제거한 문자열로 변환"""
    if value is None:
        return ''

    return str(value).strip()


def _number(value, header):
    """엑셀 숫자 열을 정수로 변환하고 잘못된 값은 명확한 오류로 처리"""
    try:
        return int(float(value))

    except (TypeError, ValueError) as exc:
        raise CommandError(
            f'{header} 열에 숫자가 아닌 값이 있음: {value!r}'
        ) from exc


def read_excel_rows():
    """students.xlsx를 읽어 StudentScore 모델에 저장할 딕셔너리 목록을 반환"""

    # students.xlsx가 프로젝트 최상위 폴더에 있는지 확인
    if not EXCEL_FILE.exists():
        raise CommandError(
            f'{EXCEL_FILE} 파일이 없음. '
            'manage.py와 같은 위치에 students.xlsx를 배치'
        )

    try:
        # 엑셀 파일을 읽기 전용으로 열어 메모리 사용량을 줄임
        workbook = load_workbook(
            EXCEL_FILE,
            read_only=True,
            data_only=True
        )

    except (InvalidFileException, OSError, ValueError) as exc:
        raise CommandError(
            'students.xlsx를 열 수 없음. '
            '올바른 xlsx 파일인지 확인'
        ) from exc

    try:
        # data 시트를 우선 사용
        # data 시트가 없으면 첫 번째 시트를 사용
        if 'data' in workbook.sheetnames:
            worksheet = workbook['data']
        else:
            worksheet = workbook.worksheets[0]

        # 엑셀 시트의 모든 행을 읽음
        rows = worksheet.iter_rows(
            values_only=True
        )

        try:
            # 첫 번째 행은 엑셀 헤더
            raw_headers = next(rows)

        except StopIteration as exc:
            raise CommandError(
                '엑셀 시트에 헤더 행이 없음'
            ) from exc

        # 헤더의 공백을 제거하고 문자열로 변환
        headers = [
            _text(value)
            for value in raw_headers
        ]

        # 엑셀 헤더 이름과 실제 열 번호 연결
        indexes = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        # 필요한 헤더가 모두 있는지 확인
        missing = [
            header
            for header in COLUMN_ALIASES
            if header not in indexes
        ]

        if missing:
            raise CommandError(
                '엑셀 헤더가 누락됨: '
                + ', '.join(missing)
            )

        # DB에 저장할 학생 목록
        result = []

        # 헤더 다음 행부터 학생 데이터 읽기
        for raw_row in rows:

            # 완전히 비어 있는 엑셀 행은 건너뜀
            if not any(
                value is not None and str(value).strip()
                for value in raw_row
            ):
                continue

            # 한 명의 학생 데이터
            values = {}

            # 엑셀 열을 모델 필드에 맞게 변환
            for excel_header, model_field in COLUMN_ALIASES.items():

                # 현재 헤더의 열 번호
                index = indexes[excel_header]

                # 행에 해당 열이 없으면 빈 문자열 사용
                if index < len(raw_row):
                    value = raw_row[index]
                else:
                    value = ''

                # 숫자로 저장할 필드
                if model_field in {
                    'no',
                    'grade',
                    'stature',
                    'kor',
                    'eng',
                    'math',
                }:
                    values[model_field] = _number(
                        value,
                        excel_header
                    )

                # 문자로 저장할 필드
                else:
                    values[model_field] = _text(value)

            # 변환이 끝난 학생 데이터를 목록에 추가
            result.append(values)

        return result

    finally:
        # 엑셀 파일 핸들을 닫음
        workbook.close()


class Command(BaseCommand):
    """python manage.py import_students 실행 처리"""

    help = 'students.xlsx 데이터를 관리자용 StudentScore 모델로 가져옴'

    def handle(self, *args, **options):

        # StudentScore 모델이 사용하는 실제 DB 테이블 이름
        table_name = StudentScore._meta.db_table

        # Django 마이그레이션 기록이 아니라
        # 실제 DB에 테이블이 존재하는지 확인
        existing_tables = connection.introspection.table_names()

        # 테이블이 실제로 없으면 직접 생성
        #
        # 오류 상황:
        # django.db.utils.OperationalError:
        # no such table: stuscore_studentscore
        #
        # 위 오류를 막기 위한 처리
        if table_name not in existing_tables:
            with connection.schema_editor() as schema_editor:
                schema_editor.create_model(StudentScore)

        # students.xlsx에서 학생 데이터를 읽음
        students = read_excel_rows()

        # 새로 생성된 학생 수
        created_count = 0

        # 기존 데이터가 갱신된 학생 수
        updated_count = 0

        # 학생 데이터를 한 명씩 DB에 저장
        for values in students:

            # 번호가 같으면 기존 학생 정보를 갱신
            # 번호가 없으면 새 학생 정보를 생성
            _, created = StudentScore.objects.update_or_create(
                no=values['no'],
                defaults=values
            )

            if created:
                created_count += 1
            else:
                updated_count += 1

        # 실행 결과 출력
        self.stdout.write(
            self.style.SUCCESS(
                f'완료: 신규 {created_count}명, '
                f'갱신 {updated_count}명, '
                f'전체 {len(students)}명'
            )
        )
