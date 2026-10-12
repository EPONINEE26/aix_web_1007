
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from students.models import Stu


# manage.py와 students.xlsx가 함께 있는 프로젝트 최상위 폴더
EXCEL_FILE = Path(__file__).resolve().parents[3] / 'students.xlsx'


def text_value(value):
    # 빈 셀은 빈 문자열로 변환
    if value is None:
        return ''

    # 셀 값을 문자열로 변환하고 앞뒤 공백 제거
    return str(value).strip()


def number_value(value, header):
    # 빈 숫자 셀은 0으로 변환
    if value is None or value == '':
        return 0

    try:
        # 엑셀 숫자 값을 정수로 변환
        return int(float(value))

    except (TypeError, ValueError) as exc:
        raise CommandError(
            f'{header} 열의 값이 숫자가 아님: {value!r}'
        ) from exc


def read_students_excel():
    # students.xlsx 파일이 존재하는지 확인
    if not EXCEL_FILE.exists():
        raise CommandError(
            f'students.xlsx 파일을 찾을 수 없음: {EXCEL_FILE}'
        )

    try:
        # 엑셀 파일을 읽기 전용으로 열기
        workbook = load_workbook(
            EXCEL_FILE,
            read_only=True,
            data_only=True
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
    ) as exc:
        raise CommandError(
            'students.xlsx 파일을 열 수 없음'
        ) from exc

    try:
        # data 시트가 있으면 data 시트 사용
        # 없으면 첫 번째 시트 사용
        if 'data' in workbook.sheetnames:
            worksheet = workbook['data']
        else:
            worksheet = workbook.worksheets[0]

        # 엑셀 전체 행 읽기
        rows = worksheet.iter_rows(values_only=True)

        try:
            # 첫 번째 행은 헤더
            raw_headers = next(rows)

        except StopIteration as exc:
            raise CommandError(
                '엑셀 시트에 헤더가 없음'
            ) from exc

        # 헤더 공백 제거
        headers = [
            text_value(value)
            for value in raw_headers
        ]

        # 헤더 이름과 열 번호 연결
        header_indexes = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        # 반드시 필요한 엑셀 헤더
        required_headers = [
            '이름',
            '학교',
            '학년',
            '키',
            'SW특기',
        ]

        # 누락된 헤더 확인
        missing_headers = [
            header
            for header in required_headers
            if header not in header_indexes
        ]

        if missing_headers:
            raise CommandError(
                '엑셀 헤더가 누락됨: '
                + ', '.join(missing_headers)
            )

        # 저장할 학생 목록
        students = []

        # 두 번째 행부터 학생 데이터 읽기
        for raw_row in rows:

            # 완전히 빈 행은 건너뜀
            if not any(
                value is not None and str(value).strip()
                for value in raw_row
            ):
                continue

            # 각 열 번호
            name_index = header_indexes['이름']
            school_index = header_indexes['학교']
            grade_index = header_indexes['학년']
            stature_index = header_indexes['키']
            sw_index = header_indexes['SW특기']

            # 셀 값 읽기
            name = text_value(raw_row[name_index])
            school = text_value(raw_row[school_index])
            grade = number_value(
                raw_row[grade_index],
                '학년'
            )
            stature = number_value(
                raw_row[stature_index],
                '키'
            )
            sw = text_value(raw_row[sw_index])

            # 이름이 없는 행은 건너뜀
            if not name:
                continue

            # 기존 Stu 모델 필드 이름으로 변환
            students.append(
                {
                    'name': name,
                    'major': school,
                    'age': stature,
                    'grade': grade,
                    'gender': sw,
                }
            )

        return students

    finally:
        # 엑셀 파일 닫기
        workbook.close()


class Command(BaseCommand):
    # 관리 명령 설명
    help = 'students.xlsx 데이터를 기존 Stu 모델에 저장'

    def handle(self, *args, **options):
        # 엑셀에서 학생 데이터 읽기
        students = read_students_excel()

        # 신규 생성 수
        created_count = 0

        # 기존 데이터 수정 수
        updated_count = 0

        # 학생 데이터를 한 명씩 DB에 저장
        for student in students:

            # 이름이 같으면 기존 학생 수정
            # 이름이 없으면 새 학생 생성
            _, created = Stu.objects.update_or_create(
                name=student['name'],
                defaults={
                    'major': student['major'],
                    'age': student['age'],
                    'grade': student['grade'],
                    'gender': student['gender'],
                }
            )

            if created:
                created_count += 1
            else:
                updated_count += 1

        # 처리 결과 출력
        self.stdout.write(
            self.style.SUCCESS(
                f'완료: 신규 {created_count}명, '
                f'갱신 {updated_count}명, '
                f'전체 {len(students)}명'
            )
        )
