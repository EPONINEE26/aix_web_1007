# 파일 경로 계산에 사용하는 Path.
from pathlib import Path

# Django 터미널 명령을 만들 때 상속하는 기본 클래스와 사용자 오류 클래스.
from django.core.management.base import BaseCommand, CommandError

# xlsx 파일을 읽는 함수와 잘못된 엑셀 파일 오류 클래스.
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

# 읽은 값을 저장할 기존 학생 모델.
from students.models import Stu


# 현재 파일 위치에서 프로젝트 최상위까지 올라가 students.xlsx 위치 계산.
EXCEL_FILE = Path(__file__).resolve().parents[3] / 'students.xlsx'


def text_value(value):
    # 엑셀 빈 셀(None)을 DB에 넣을 빈 문자열로 바꾸는 보조 함수.
    if value is None:
        return ''

    # 숫자나 날짜 등 다른 셀 값도 문자열로 통일하고 양쪽 공백 제거.
    return str(value).strip()


def number_value(value, header):
    # 숫자 셀이 비어 있으면 Stu 모델의 기본값과 같은 0으로 처리.
    if value is None or value == '':
        return 0

    try:
        # 3.0처럼 엑셀에서 실수 형태로 읽힌 숫자도 정수 3으로 변환.
        return int(float(value))
    except (TypeError, ValueError) as exc:
        # 어느 열에서 어떤 값이 실패했는지 CommandError로 사용자에게 전달.
        raise CommandError(
            f'{header} 열의 값이 숫자가 아님: {value!r}'
        ) from exc


def read_students_excel():
    # 파일이 없으면 openpyxl을 호출하기 전에 명확한 오류를 발생.
    if not EXCEL_FILE.exists():
        raise CommandError(f'students.xlsx 파일을 찾을 수 없음: {EXCEL_FILE}')

    try:
        # 읽기 전용·계산 결과 모드로 엑셀 열기.
        workbook = load_workbook(
            EXCEL_FILE,
            read_only=True,
            data_only=True,
        )
    except (InvalidFileException, OSError, ValueError) as exc:
        raise CommandError('students.xlsx 파일을 열 수 없음') from exc

    try:
        # data 시트가 있으면 그 시트, 없으면 첫 번째 시트 선택.
        worksheet = (
            workbook['data']
            if 'data' in workbook.sheetnames
            else workbook.worksheets[0]
        )

        # 각 행을 셀 값 튜플로 읽는 반복자.
        rows = worksheet.iter_rows(values_only=True)

        try:
            # 첫 행은 학생 데이터가 아니라 열 제목으로 사용.
            raw_headers = next(rows)
        except StopIteration as exc:
            raise CommandError('엑셀 시트에 헤더가 없음') from exc

        # 헤더의 빈 값과 공백을 정리해 비교 가능한 문자열 목록으로 변환.
        headers = [text_value(value) for value in raw_headers]

        # {'이름': 0, '학교': 1}처럼 헤더명과 실제 열 번호를 연결.
        header_indexes = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        # 이 명령이 Stu에 저장하려면 반드시 필요한 엑셀 열.
        required_headers = ['이름', '학교', '학년', '키', 'SW특기']
        missing_headers = [
            header
            for header in required_headers
            if header not in header_indexes
        ]
        if missing_headers:
            raise CommandError(
                '엑셀 헤더가 누락됨: ' + ', '.join(missing_headers)
            )

        students = []

        # 헤더 다음 행부터 학생 데이터를 한 행씩 변환.
        for raw_row in rows:
            # 완전히 빈 행은 학생 한 명으로 만들지 않고 건너뜀.
            if not any(
                value is not None and str(value).strip()
                for value in raw_row
            ):
                continue

            # 헤더 이름으로 열 위치를 찾은 뒤 현재 행의 셀 값 읽기.
            name = text_value(raw_row[header_indexes['이름']])
            school = text_value(raw_row[header_indexes['학교']])
            grade = number_value(
                raw_row[header_indexes['학년']],
                '학년',
            )
            stature = number_value(
                raw_row[header_indexes['키']],
                '키',
            )
            sw = text_value(raw_row[header_indexes['SW특기']])

            # 이름이 없으면 update_or_create의 기준값으로 사용할 수 없어 제외.
            if not name:
                continue

            # 엑셀 이름을 Stu 모델 필드 이름으로 바꾼 딕셔너리.
            students.append({
                'name': name,
                'major': school,
                'age': stature,
                'grade': grade,
                'gender': sw,
            })

        return students

    finally:
        # 변환 성공 여부와 관계없이 열린 엑셀 파일 닫기.
        workbook.close()


# python manage.py import_students_excel 명령의 실행 클래스.
class Command(BaseCommand):
    # python manage.py help에서 보여줄 명령 설명.
    help = 'students.xlsx 데이터를 기존 Stu 모델에 저장'

    def handle(self, *args, **options):
        # 엑셀을 모델 필드 딕셔너리 목록으로 변환.
        students = read_students_excel()
        created_count = 0
        updated_count = 0

        for student in students:
            # name이 같은 학생을 찾고, 있으면 defaults로 수정.
            # 없으면 name과 defaults를 합쳐 새 Stu 객체 생성.
            _, created = Stu.objects.update_or_create(
                name=student['name'],
                defaults={
                    'major': student['major'],
                    'age': student['age'],
                    'grade': student['grade'],
                    'gender': student['gender'],
                },
            )

            # created=True면 새 행, False면 기존 행 수정.
            if created:
                created_count += 1
            else:
                updated_count += 1

        # 명령 실행 결과를 터미널에 성공 메시지로 출력.
        self.stdout.write(
            self.style.SUCCESS(
                f'완료: 신규 {created_count}명, '
                f'갱신 {updated_count}명, '
                f'전체 {len(students)}명'
            )
        )
