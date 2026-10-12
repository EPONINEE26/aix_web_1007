from pathlib import Path

from django.conf import settings
from django.shortcuts import redirect, render
from openpyxl import load_workbook

from students.models import Stu


def swrite(request):
    """
    학생 기본 정보 입력 화면과 저장 요청을 모두 처리하는 뷰.
    브라우저가 화면을 열 때와 저장 버튼을 눌렀을 때 같은 주소로 요청이 오지만,
    요청 방식이 GET인지 POST인지에 따라 실행할 작업이 달라짐.
    """

    if request.method == 'GET':
        # render(request, 템플릿 이름): HTML 파일을 찾아 브라우저에 반환.
        # GET에서는 저장할 입력값이 없으므로 입력 화면만 반환.
        return render(request, 'swrite.html')

    if request.method == 'POST':
        name = request.POST.get('name')
        major = request.POST.get('major')
        grade = request.POST.get('grade')
        age = request.POST.get('age')
        gender = request.POST.get('gender')

        # Stu(...): 아직 DB에 저장되지 않은 학생 객체를 메모리에 생성.
        # 각 키워드의 왼쪽 이름은 모델 필드, 오른쪽 값은 입력받은 내용.
        student = Stu(
            name=name,
            major=major,
            grade=grade,
            age=age,
            gender=gender,
        )

        # save(): 위에서 만든 객체를 실제 데이터베이스 행으로 저장.
        student.save()

        # redirect('/'): 저장 완료 후 메인 페이지로 새 요청을 보내는 처리.
        # POST 결과 화면에서 새로고침할 때 같은 데이터가 다시 저장되는 문제를 줄이는 흐름.
        return redirect('/')


def slist(request):
    """
    프로젝트 최상위의 students.xlsx를 읽어 표 화면으로 전달하는 뷰.
    이 함수는 Stu 데이터베이스를 조회하지 않고 엑셀 파일을 직접 읽음.
    """

    # settings.BASE_DIR: manage.py가 있는 프로젝트 최상위 경로.
    # Path / '파일명': 운영체제에 맞는 파일 경로를 조합하는 표현.
    excel_path = Path(settings.BASE_DIR) / 'students.xlsx'

    # exists(): 해당 경로에 실제 파일이 있는지 확인.
    # 파일이 없으면 엑셀을 열 수 없으므로 템플릿에 빈 값과 오류 문구 전달.
    if not excel_path.exists():
        return render(request, 'slist.html', {
            'headers': [],
            'rows': [],
            'error': 'students.xlsx 파일을 프로젝트 최상위 폴더에 넣어야 함.',
        })

    # load_workbook(): openpyxl로 xlsx 파일 열기.
    # read_only=True: 읽기만 하는 파일을 메모리 부담이 적은 방식으로 열기.
    # data_only=True: 수식 자체가 아니라 엑셀에 저장된 계산 결과 읽기.
    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
    )

    try:
        # workbook.active: 엑셀에서 현재 활성화된 시트.
        sheet = workbook.active

        # iter_rows(): 시트를 행 단위로 순회.
        # values_only=True: 셀 객체가 아니라 셀 안의 실제 값만 가져오기.
        # list(): 반복 결과를 목록으로 바꿔 첫 행과 나머지 행을 분리할 수 있게 함.
        data = list(sheet.iter_rows(values_only=True))

        # data[0]: 첫 번째 엑셀 행.
        # 이 프로젝트에서는 첫 행을 표의 열 제목으로 사용.
        # 엑셀이 완전히 비어 있으면 data가 빈 목록이므로 headers도 빈 목록.
        headers = list(data[0]) if data else []

        # data[1:]: 첫 번째 행을 제외한 학생 데이터.
        # any(...): 한 행 안에 값이 하나라도 있는지 확인.
        # 완전히 빈 행은 화면에서 제외.
        rows = [
            list(row)
            for row in data[1:]
            if any(value is not None for value in row)
        ]

    finally:
        # finally: 정상 처리와 오류 발생 여부에 관계없이 실행되는 영역.
        # 엑셀 파일을 닫지 않으면 파일 잠금이나 리소스 누수가 생길 수 있음.
        workbook.close()

    # 세 번째 인자는 템플릿에 전달할 context 딕셔너리.
    # 템플릿에서는 같은 이름인 headers, rows, error로 값에 접근.
    return render(request, 'slist.html', {
        'headers': headers,
        'rows': rows,
        'error': '',
    })
